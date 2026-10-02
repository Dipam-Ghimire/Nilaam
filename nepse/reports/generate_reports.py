"""Generate project_report.docx and presentation.pptx."""
import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from pptx import Presentation
from pptx.util import Inches as PInches, Pt as PPt

OUT = os.path.dirname(__file__)

# ---------- DOCX ----------
doc = Document()
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)

doc.add_heading("NEPSE Stock Price Prediction System", 0)
doc.add_paragraph("Using Machine Learning and Data Science").italic = True

sections = [
    ("1. Introduction",
     "The Nepal Stock Exchange (NEPSE) is the only stock exchange of Nepal. "
     "Predicting stock prices supports investors and analysts in decision-making. "
     "This project builds an end-to-end pipeline that collects NEPSE data, "
     "engineers technical features, trains classical and deep-learning models, "
     "and serves an interactive Streamlit dashboard."),
    ("2. Objectives",
     "• Collect historical OHLCV data for major NEPSE companies.\n"
     "• Perform EDA and feature engineering (MA, EMA, RSI, MACD, lags, volatility).\n"
     "• Train Linear Regression, Random Forest, XGBoost, and LSTM models.\n"
     "• Compare classical ML vs deep learning on RMSE / MAE / R² / MAPE.\n"
     "• Deliver an interactive dashboard for analytics and forecasting."),
    ("3. Data Collection",
     "Live NEPSE data is fetched using the nepse-scraper Python package. "
     "If the live endpoint is unavailable, a realistic synthetic OHLCV generator "
     "produces drop-in data so the rest of the pipeline runs end-to-end. "
     "Companies covered: Nabil Bank, NIC Asia Bank, Nepal Reinsurance, "
     "Citizen Investment Trust, Shivam Cements, HIDCL."),
    ("4. Preprocessing & Feature Engineering",
     "Data is sorted by date, de-duplicated, numerically coerced, and "
     "forward/back-filled. Engineered features include moving averages (5/10/20/50), "
     "exponential MAs, RSI(14), MACD with signal & histogram, daily returns, "
     "rolling volatility, lag features (1–10), high-low and open-close spreads, "
     "and 10-day volume MA. The target is the next-day closing price."),
    ("5. Models",
     "• Linear Regression — baseline.\n"
     "• Random Forest Regressor — non-linear ensemble.\n"
     "• XGBoost Regressor — gradient-boosted trees (typically best classical model).\n"
     "• LSTM — recurrent neural network on a 30-day sliding window of scaled closes."),
    ("6. Evaluation",
     "Models are evaluated on a chronological 80/20 split using RMSE, MAE, R², "
     "and MAPE. Metrics are persisted to models/<SYMBOL>_metrics.json. "
     "On typical NEPSE series, XGBoost and LSTM tend to outperform Linear Regression, "
     "with LSTM capturing longer-term momentum and XGBoost excelling at short horizons."),
    ("7. Dashboard",
     "A multi-page Streamlit app provides: (a) Market Overview with candlestick "
     "charts and KPIs, (b) Technical Analysis with MA/RSI/MACD overlays, "
     "(c) Forecast view for 7/15/30-day horizons, and (d) one-click model retraining."),
    ("8. How to Run",
     "1. pip install -r requirements.txt\n"
     "2. python src/train_model.py --symbol NABIL\n"
     "3. streamlit run app/streamlit_app.py"),
    ("9. Conclusion",
     "The system demonstrates a complete data-science workflow on the Nepalese "
     "stock market: ingestion, feature engineering, modelling, evaluation, and "
     "deployment. It provides investors with data-driven insights and a "
     "reusable framework for further research."),
]
for title, body in sections:
    doc.add_heading(title, level=1)
    for para in body.split("\n"):
        doc.add_paragraph(para)

doc.save(os.path.join(OUT, "project_report.docx"))

# ---------- PPTX ----------
prs = Presentation()
prs.slide_width = PInches(13.333)
prs.slide_height = PInches(7.5)

def add_slide(title, bullets):
    layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = title
    tf = slide.placeholders[1].text_frame
    tf.text = bullets[0]
    for b in bullets[1:]:
        p = tf.add_paragraph()
        p.text = b
        p.font.size = PPt(18)

# Title slide
title_slide = prs.slides.add_slide(prs.slide_layouts[0])
title_slide.shapes.title.text = "NEPSE Stock Price Predictor"
title_slide.placeholders[1].text = "Machine Learning & Data Science on the Nepal Stock Exchange"

add_slide("Problem Statement", [
    "NEPSE investors lack accessible, data-driven price-prediction tools.",
    "Goal: forecast next-day & multi-day closing prices for major companies.",
    "Compare classical ML and deep learning approaches.",
])
add_slide("Companies Covered", [
    "Nabil Bank (NABIL)",
    "NIC Asia Bank (NICA)",
    "Nepal Reinsurance Company (NRIC)",
    "Citizen Investment Trust (CIT)",
    "Shivam Cements (SHIVM)",
    "Hydroelectricity Investment & Development Co. (HIDCL)",
])
add_slide("Data Pipeline", [
    "Live data via nepse-scraper (synthetic fallback).",
    "Cleaning: sort, de-duplicate, ffill/bfill.",
    "Features: MA, EMA, RSI(14), MACD, lags, volatility, spreads.",
    "Target: next-day close.",
])
add_slide("Models", [
    "Linear Regression — baseline.",
    "Random Forest — non-linear ensemble.",
    "XGBoost — gradient boosting (typically best classical).",
    "LSTM — 30-day sliding window on scaled closes.",
])
add_slide("Evaluation", [
    "Chronological 80/20 split (no look-ahead).",
    "Metrics: RMSE, MAE, R², MAPE.",
    "Persisted to models/<SYMBOL>_metrics.json.",
])
add_slide("Dashboard (Streamlit)", [
    "Market Overview: candlestick + KPIs.",
    "Technical Analysis: MA, RSI, MACD.",
    "Forecast: 7 / 15 / 30 days.",
    "One-click retraining per symbol.",
])
add_slide("How to Run", [
    "pip install -r requirements.txt",
    "python src/train_model.py --symbol NABIL",
    "streamlit run app/streamlit_app.py",
])
add_slide("Conclusion", [
    "End-to-end NEPSE prediction pipeline delivered.",
    "Classical ML vs DL comparison enabled.",
    "Reusable framework for further research.",
])

prs.save(os.path.join(OUT, "presentation.pptx"))
print("Reports generated.")
